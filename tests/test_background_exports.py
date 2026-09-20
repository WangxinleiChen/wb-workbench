"""Export checks use explicit numerical fixtures, never live experiment data."""
import copy
import csv
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

import numpy as np
from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import exports


class FixtureStore:
    def __init__(self, directory, experiment, result):
        self.directory = directory
        self.experiment = experiment
        self.result = result
        self.state = {"audit": [{"experimentId": experiment["id"], "action": "background.apply"}]}

    def results(self, exp):
        return self.result

    def image_path(self, image):
        return self.directory / "images" / image["id"] / "original"

    def annotated(self, exp, role):
        return self.image_path(exp["images"][role]).read_bytes()


class BackgroundExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.exp = {"id": "exp-test", "name": "Export fixture", "revision": 7, "notes": "",
                    "settings": {r: {"polarity": "dark"} for r in ("pho", "total")},
                    "images": {}, "background": {}, "samples": [{"id": "s1", "name": "Control"}]}
        row = {"recordId": "exp-test:s1", "sampleId": "s1", "sampleName": "Control", "group": "Ctr",
               "control": True, "status": "已确认", "warnings": [], "ratio": 0.5, "relative": 1.0}
        for role, raw_sum, net in (("pho", 600.0, 200.0), ("total", 400.0, 400.0)):
            original = self.directory / "images" / role / "original"
            original.parent.mkdir(parents=True)
            Image.fromarray(np.full((8, 8), 200, dtype=np.uint8)).save(original, format="PNG")
            sha = hashlib.sha256(original.read_bytes()).hexdigest()
            metadata = {"id": role, "filename": role + ".png", "sha256": sha, "bitDepth": 8, "warnings": []}
            self.exp["images"][role] = metadata
            row[role + "Image"] = metadata
            row[role + "Roi"] = {"confirmed": True, "band": {"x": 2, "y": 2, "w": 2, "h": 2},
                                  "background": {"x": 2, "y": 0, "w": 2, "h": 1}}
            row[role] = {"rawSum": raw_sum, "area": 4, "net": net, "backgroundMean": 99.0,
                         "backgroundArea": 0, "endpointCount": 0, "warnings": [], "valid": True,
                         "method": "model", "backgroundSource": "robust-quadratic-v1", "backgroundContribution": 800.0}
            artifact = self.artifact(role, sha, 2.5)
            self.exp["background"][role] = {"mode": "model", "applied": artifact, "preview": artifact}
        self.result = {"rows": [row], "controlMean": 0.5, "controlReady": True, "warnings": [],
                       "method": "模式 A 局部背景均值；模式 B 一次模型背景扣除"}
        self.store = FixtureStore(self.directory, self.exp, self.result)

    def tearDown(self):
        self.temp.cleanup()

    def artifact(self, suffix, sha, sigma):
        key = hashlib.sha256(suffix.encode()).hexdigest()
        folder = self.directory / "derived" / key
        folder.mkdir(parents=True, exist_ok=True)
        metadata = {"key": key, "algorithm": "robust-quadratic-v1", "algorithmVersion": 1,
                    "params": {"clipSigma": sigma}, "sourceSha256": sha, "polarity": "dark",
                    "region": {"x": 0, "y": 0, "w": 8, "h": 8}, "display": {"purpose": "display-only"}}
        (folder / "manifest.json").write_text(json.dumps(metadata))
        np.save(folder / "background.npy", np.full((8, 8), 200.0, dtype=np.float64))
        signal = np.zeros((8, 8), dtype=np.float64)
        signal[0, 0] = -2.0
        np.save(folder / "corrected.npy", signal)
        for name in ("background.png", "corrected.png"):
            Image.fromarray(np.zeros((8, 8), dtype=np.uint8)).save(folder / name)
        return metadata

    def xlsx(self):
        data, _, _ = exports.build_export(self.store, [self.exp], "xlsx")
        return load_workbook(io.BytesIO(data), data_only=False), load_workbook(io.BytesIO(data), data_only=True)

    def test_model_formula_uses_one_background_contribution_and_cached_values(self):
        # The deliberately different legacy background mean (99) must not enter this formula.
        formulas, cached = self.xlsx()
        detail = formulas["测量明细"]
        self.assertEqual(detail["P2"].value, '=IF(H2="dark",V2-M2,M2-V2)')
        self.assertNotIn("N2", detail["P2"].value)
        self.assertEqual(cached["测量明细"]["P2"].value, 200.0)
        self.assertEqual(cached["测量明细"]["V2"].value, 800.0)
        self.assertEqual(cached["定量结果"]["J2"].value, 0.5)
        self.assertEqual(cached["定量结果"]["O2"].value, "model")
        self.assertEqual(cached["测量明细"]["T2"].value, "model")
        self.assertEqual(cached["测量明细"]["Y2"].value, '{"clipSigma":2.5}')
        # Existing columns remain A:S and retain their meaning.
        self.assertEqual(detail["S1"].value, "原图在归档内的位置")
        self.assertEqual(detail["T1"].value, "定量模式")
        self.assertEqual(cached["方法与来源"]["B2"].value, "WB Workbench v1.1.0-dev")
        method_sheet = {row[0]: row[1] for row in cached["方法与来源"].values}
        self.assertIn("模式 B（model）", method_sheet["测量定义"])
        self.assertIn("sum(B−I)", method_sheet["测量定义"])
        self.assertIn("模式 B 不再减局部背景框", method_sheet["测量定义"])
        self.assertIn("P=模式 B（model）", method_sheet["实验计算方法"])

    def test_bright_formula_and_negative_net_not_clipped_and_publish_gate_retained(self):
        row = self.result["rows"][0]
        self.exp["settings"]["pho"]["polarity"] = "bright"
        self.exp["background"]["pho"]["applied"]["polarity"] = "bright"
        row["pho"].update(rawSum=780.0, net=-20.0, valid=False)
        row["ratio"] = row["relative"] = self.result["controlMean"] = None
        row["status"] = "无效信号"
        formulas, cached = self.xlsx()
        self.assertEqual(cached["测量明细"]["P2"].value, -20.0)
        self.assertIn("M2-V2", formulas["测量明细"]["P2"].value)
        self.assertEqual(cached["定量结果"]["N2"].value, 0)
        self.assertIsNone(cached["定量结果"]["J2"].value)
        self.assertIsNone(cached["定量结果"]["L2"].value)
        self.assertIn("N2=1", formulas["定量结果"]["J2"].value)

    def test_old_local_export_formula_and_fallback_contribution_unchanged(self):
        self.exp.pop("background")
        row = self.result["rows"][0]
        for role in ("pho", "total"):
            row[role].pop("method")
            row[role].pop("backgroundSource")
            row[role].pop("backgroundContribution")
            row[role]["backgroundMean"] = 200.0
            row[role]["backgroundArea"] = 2
        formulas, cached = self.xlsx()
        self.assertEqual(formulas["测量明细"]["P2"].value, '=IF(H2="dark",L2*N2-M2,M2-L2*N2)')
        self.assertEqual(cached["测量明细"]["P2"].value, 200.0)
        self.assertEqual(cached["测量明细"]["V2"].value, 800.0)
        self.assertEqual(cached["测量明细"]["T2"].value, "local")

    def test_preview_params_separate_from_formal_model_in_csv_and_workbook(self):
        preview = self.artifact("pho-preview", self.exp["images"]["pho"]["sha256"], 4.0)
        self.exp["background"]["pho"]["preview"] = preview
        self.result["rows"][0]["sampleName"] = "=2+3"
        payload, _, _ = exports.build_export(self.store, [self.exp], "csv")
        table = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
        row = table[0]
        self.assertEqual(row["样本名称"], "'=2+3")
        self.assertEqual(row["P_正式方法参数"], '{"clipSigma":2.5}')
        self.assertEqual(row["P_预览参数"], '{"clipSigma":4.0}')
        self.assertEqual(row["P_背景贡献"], "800.0")
        self.assertEqual(row["P_方法状态"], "正式应用")
        self.assertEqual(float(row["p/total图像信号比"]), 0.5)
        self.assertEqual(row["P_方法原图SHA256"], self.exp["images"]["pho"]["sha256"])
        self.assertEqual(row["P_预览派生结果键"], preview["key"])
        self.assertIn("P=模式 B（model）", row["计算方法"])
        self.assertIn("total=模式 B（model）", row["计算方法"])
        self.assertIn("sum(I−B)", row["计算方法"])
        self.assertIn("模式 B 不再减局部背景框", row["计算方法"])
        _, cached = self.xlsx()
        self.assertEqual(cached["测量明细"]["AE2"].value, row["P_预览参数"])
        self.assertEqual(cached["定量结果"]["J2"].value, 0.5)

    def test_invalidated_model_is_blank_instead_of_falling_back_to_local(self):
        self.exp["background"]["pho"]["applied"] = None
        row = self.result["rows"][0]
        row["pho"] = None
        row["ratio"] = row["relative"] = self.result["controlMean"] = None
        row["phoRoi"]["confirmed"] = False
        formulas, cached = self.xlsx()
        self.assertIsNone(formulas["测量明细"]["P2"].value)
        self.assertIsNone(cached["测量明细"]["V2"].value)
        self.assertIsNone(cached["定量结果"]["H2"].value)
        self.assertIsNone(cached["定量结果"]["J2"].value)
        self.assertEqual(cached["测量明细"]["T2"].value, "model")
        self.assertIn("失效", cached["测量明细"]["AC2"].value)

    def test_method_description_reports_formal_mixed_roles_not_preview_modes(self):
        # A total-image model preview exists, but the formal total method is still local.
        self.exp["background"]["total"].update(mode="local", applied=None)
        text = exports.method_description(self.exp, self.result)
        self.assertIn("P=模式 B（model）；total=模式 A（local）", text)
        self.assertIn("模式 B 不再减局部背景框", text)
        self.assertNotIn("total=模式 B（model）", text)

    def test_zip_contains_unique_applied_and_preview_artifacts_and_exact_originals(self):
        preview = self.artifact("pho-preview", self.exp["images"]["pho"]["sha256"], 4.0)
        self.exp["background"]["pho"]["preview"] = preview
        payload, _, _ = exports.build_export(self.store, [self.exp], "zip")
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), len(set(names)))
            index = json.loads(archive.read("derived/index.json"))
            self.assertEqual(len(index["usage"]), 4)
            self.assertEqual(sum(r["active"] for r in index["usage"]), 2)
            for role in ("pho", "total"):
                image = self.exp["images"][role]
                self.assertEqual(archive.read(f"originals/{image['id']}/{image['filename']}"), self.store.image_path(image).read_bytes())
            for usage in index["usage"]:
                key = usage["key"]
                for name in exports.DERIVED_FILES:
                    expected = (self.directory / "derived" / key / name).read_bytes()
                    self.assertEqual(archive.read(f"derived/{key}/{name}"), expected)
                corrected = np.load(io.BytesIO(archive.read(f"derived/{key}/corrected.npy")), allow_pickle=False)
                self.assertEqual(corrected.dtype, np.dtype("float64"))
                self.assertEqual(corrected[0, 0], -2.0)
            snapshot = json.loads(archive.read("provenance.json"))
            self.assertEqual(snapshot["experiments"][0]["background"], self.exp["background"])
            self.assertIn("不是仪器原图", archive.read("README.txt").decode())

    def test_zip_rejects_missing_derived_files_and_invalid_key(self):
        artifact = self.exp["background"]["pho"]["applied"]
        path = self.directory / "derived" / artifact["key"] / "corrected.npy"
        original = path.read_bytes()
        path.unlink()
        with self.assertRaisesRegex(ValueError, "缺少文件"):
            exports.build_export(self.store, [self.exp], "zip")
        path.write_bytes(original)
        artifact["key"] = "../outside"
        with self.assertRaisesRegex(ValueError, "缓存键无效"):
            exports.build_export(self.store, [self.exp], "zip")


if __name__ == "__main__":
    unittest.main()
