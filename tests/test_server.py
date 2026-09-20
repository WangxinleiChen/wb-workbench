"""Independent persistence, release-gating, and export regression tests.

All application data lives in a TemporaryDirectory. Synthetic measurements have
known expected ratios, and WBTEST is only a source/geometry workflow fixture.
"""

import copy
import csv
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen
import zipfile

import numpy as np
from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import exports
from server import Conflict, Handler, Store, ThreadingHTTPServer


class StoreAndExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name) / "data")

    def tearDown(self):
        self.temp.cleanup()

    def synthetic(self, name="Known experiment", pho_values=(180, 160), total_values=(160, 160), confirm=True):
        exp = self.store.create(name, 2, 1)
        for role, values in (("pho", pho_values), ("total", total_values)):
            pixels = np.full((32, 60), 200, dtype=np.uint8)
            for lane, value in enumerate(values):
                pixels[12:18, 8 + 24 * lane:16 + 24 * lane] = value
            output = io.BytesIO()
            Image.fromarray(pixels).save(output, format="PNG")
            self.store.upload(exp, role, role + ".png", output.getvalue())
        incoming = copy.deepcopy(exp)
        for role in ("pho", "total"):
            incoming["rois"][role] = [
                {"sampleId": sample["id"], "band": {"x": 8 + 24 * lane, "y": 12, "w": 8, "h": 6},
                 "background": {"x": 8 + 24 * lane, "y": 2, "w": 8, "h": 5}, "confirmed": True}
                for lane, sample in enumerate(exp["samples"])
            ]
        self.store.update(exp, incoming)
        # A first edit must not bypass review by bundling a confirmation bit.
        self.assertTrue(all(not r["confirmed"] for role in ("pho", "total") for r in exp["rois"][role]))
        if confirm:
            self.confirm(exp)
        return exp

    def confirm(self, exp):
        incoming = copy.deepcopy(exp)
        for role in ("pho", "total"):
            for roi in incoming["rois"][role]:
                roi["confirmed"] = True
        self.store.update(exp, incoming)

    def test_wbtest_starts_unconfirmed_then_controls_average_to_one(self):
        # The bundled WBTEST screenshots are the author's unpublished inputs and are
        # not distributed with the source. Skip rather than error when they are absent,
        # matching the optional-input handling in test_analysis.py.
        if not (Path(__file__).resolve().parents[1] / "examples" / "PHO.png").exists():
            self.skipTest("Optional WBTEST input not supplied")
        exp = self.store.example()
        self.assertEqual([s["name"] for s in exp["samples"]], ["Ctr1", "Ctr2", "Ctr3", "Trt1", "Trt2", "Trt3"])
        self.assertEqual(self.store.example()["id"], exp["id"])
        self.assertEqual(len(self.store.state["experiments"]), 1)
        pending = self.store.results(exp)
        self.assertTrue(all(row["ratio"] is None and row["relative"] is None for row in pending["rows"]))
        self.assertFalse(pending["controlReady"])
        self.confirm(exp)
        result = self.store.results(exp)
        self.assertTrue(result["controlReady"])
        self.assertTrue(all(row["ratio"] > 0 for row in result["rows"]))
        controls = [row["relative"] for row in result["rows"] if row["control"]]
        self.assertAlmostEqual(sum(controls) / len(controls), 1)
        self.assertTrue(any("跨膜" in s or "不同膜" in s for s in result["warnings"]))
        self.assertTrue(any("Screenshot" in s for s in result["rows"][0]["warnings"]))

    def test_edit_revokes_confirmation_and_unconfirmed_control_blocks_relatives(self):
        exp = self.synthetic()
        first = self.store.results(exp)
        self.assertEqual([r["ratio"] for r in first["rows"]], [0.5, 1.0])
        self.assertEqual([r["relative"] for r in first["rows"]], [1.0, 2.0])
        incoming = copy.deepcopy(exp)
        incoming["rois"]["pho"][0]["band"]["x"] += 1
        incoming["rois"]["pho"][0]["confirmed"] = True
        self.store.update(exp, incoming)
        self.assertFalse(exp["rois"]["pho"][0]["confirmed"])
        self.assertTrue(exp["rois"]["total"][0]["confirmed"])
        result = self.store.results(exp)
        self.assertIsNone(result["rows"][0]["ratio"])
        self.assertEqual(result["rows"][1]["ratio"], 1.0)
        self.assertTrue(all(r["relative"] is None for r in result["rows"]))
        self.confirm(exp)
        self.assertTrue(self.store.results(exp)["controlReady"])

    def test_no_controls_keeps_sample_ratios_but_no_relative_values(self):
        exp = self.synthetic()
        incoming = copy.deepcopy(exp)
        for sample in incoming["samples"]:
            sample["control"] = False
        self.store.update(exp, incoming)
        result = self.store.results(exp)
        self.assertEqual([r["ratio"] for r in result["rows"]], [0.5, 1.0])
        self.assertTrue(all(r["relative"] is None for r in result["rows"]))
        self.assertIsNone(result["controlMean"])
        self.assertFalse(result["controlReady"])

    def test_multiple_experiments_have_independent_control_baselines(self):
        first = self.synthetic("Experiment A")
        second = self.synthetic("Experiment B", pho_values=(120, 80))
        a, b = self.store.results(first), self.store.results(second)
        self.assertEqual(a["controlMean"], 0.5)
        self.assertEqual(b["controlMean"], 2.0)
        self.assertEqual([r["relative"] for r in a["rows"]], [1.0, 2.0])
        self.assertEqual([r["relative"] for r in b["rows"]], [1.0, 1.5])
        incoming = copy.deepcopy(second)
        incoming["rois"]["pho"][0]["confirmed"] = False
        self.store.update(second, incoming)
        self.assertFalse(self.store.results(second)["controlReady"])
        self.assertEqual(self.store.results(first)["controlMean"], 0.5)
        self.assertEqual(self.store.results(first)["rows"][1]["relative"], 2.0)

    def test_revision_duplicate_ids_and_out_of_bounds_fail_atomically(self):
        exp = self.synthetic()
        snapshot = copy.deepcopy(self.store.state)
        malformed = []
        stale = copy.deepcopy(exp)
        stale["revision"] -= 1
        malformed.append((stale, Conflict))
        duplicate = copy.deepcopy(exp)
        duplicate["samples"][1]["id"] = duplicate["samples"][0]["id"]
        malformed.append((duplicate, ValueError))
        duplicate_roi = copy.deepcopy(exp)
        duplicate_roi["rois"]["total"][1]["sampleId"] = duplicate_roi["rois"]["total"][0]["sampleId"]
        malformed.append((duplicate_roi, ValueError))
        outside = copy.deepcopy(exp)
        outside["name"] = "Must not persist before validation finishes"
        outside["rois"]["total"][1]["band"]["x"] = 999
        malformed.append((outside, ValueError))
        for incoming, error in malformed:
            with self.subTest(error=error.__name__, name=incoming["name"]):
                with self.assertRaises(error):
                    self.store.update(exp, incoming)
                self.assertEqual(self.store.state, snapshot)
        loaded = Store(self.store.directory)
        self.assertEqual(loaded.state, snapshot)

    def test_confirmed_overlapping_lanes_remain_invalid(self):
        exp = self.synthetic()
        incoming = copy.deepcopy(exp)
        incoming["rois"]["pho"][1]["band"] = copy.deepcopy(incoming["rois"]["pho"][0]["band"])
        self.store.update(exp, incoming)
        self.confirm(exp)
        result = self.store.results(exp)
        self.assertTrue(all(row["confirmed"] for row in result["rows"]))
        self.assertTrue(all(row["ratio"] is None for row in result["rows"]))
        self.assertTrue(all(not row["pho"]["valid"] for row in result["rows"]))

    def test_replacing_image_clears_only_affected_role_and_preserves_prior_original(self):
        exp = self.synthetic()
        previous = copy.deepcopy(exp["images"]["pho"])
        original = self.store.image_path(previous).read_bytes()
        total_before = copy.deepcopy(exp["rois"]["total"])
        self.store.upload(exp, "pho", "new-source.png", original)
        self.assertEqual(exp["rois"]["pho"], [])
        self.assertEqual(exp["rois"]["total"], total_before)
        self.assertEqual(self.store.image_path(previous).read_bytes(), original)
        self.assertNotEqual(previous["id"], exp["images"]["pho"]["id"])
        self.assertTrue(all(r["ratio"] is None for r in self.store.results(exp)["rows"]))

    def test_restart_restores_results_sources_confirmation_and_audit(self):
        exp = self.synthetic()
        incoming = copy.deepcopy(exp)
        incoming["notes"] = "重启后仍须保留的实验说明"
        self.store.update(exp, incoming)
        before = self.store.results(exp)
        expected_revision = exp["revision"]
        reopened = Store(self.store.directory)
        restored = reopened.get(exp["id"])
        self.assertEqual(restored["notes"], incoming["notes"])
        self.assertEqual(restored["revision"], expected_revision)
        self.assertEqual(reopened.results(restored), before)
        self.assertTrue(all(roi["confirmed"] for role in ("pho", "total") for roi in restored["rois"][role]))
        for role in ("pho", "total"):
            image = restored["images"][role]
            self.assertEqual(hashlib.sha256(reopened.image_path(image).read_bytes()).hexdigest(), image["sha256"])
        recorded = [json.loads(line) for line in reopened.audit_path.read_text().splitlines()]
        self.assertEqual(recorded, reopened.state["audit"])
        self.assertEqual(recorded[-1]["revision"], expected_revision)
        payload, _, _ = exports.build_export(reopened, [restored], "csv")
        row = next(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
        self.assertEqual(float(row["p/total图像信号比"]), 0.5)

    def test_xlsx_formulas_cached_values_and_experiment_local_references(self):
        first = self.synthetic("Experiment A")
        second = self.synthetic("Experiment B", pho_values=(120, 80))
        data, mime, extension = exports.build_export(self.store, [first, second], "xlsx")
        self.assertEqual(extension, "xlsx")
        self.assertIn("spreadsheetml", mime)
        formulas = load_workbook(io.BytesIO(data), data_only=False)
        cached = load_workbook(io.BytesIO(data), data_only=True)
        self.assertEqual(formulas.sheetnames, ["定量结果", "测量明细", "方法与来源"])
        summary, cached_summary = formulas["定量结果"], cached["定量结果"]
        details, cached_details = formulas["测量明细"], cached["测量明细"]
        expected = [row for exp in (first, second) for row in self.store.results(exp)["rows"]]
        for sheet_row, result in enumerate(expected, 2):
            self.assertEqual(summary[f"A{sheet_row}"].value, result["recordId"])
            self.assertIn(f"H{sheet_row}/I{sheet_row}", summary[f"J{sheet_row}"].value)
            self.assertIn(f"N{sheet_row}=1", summary[f"J{sheet_row}"].value)
            self.assertAlmostEqual(cached_summary[f"J{sheet_row}"].value, result["ratio"])
            self.assertAlmostEqual(cached_summary[f"L{sheet_row}"].value, result["relative"])
            self.assertAlmostEqual(cached_summary[f"H{sheet_row}"].value, result["pho"]["net"])
        self.assertIn("AVERAGE(J2)", summary["K2"].value)
        self.assertIn("AVERAGE(J4)", summary["K4"].value)
        self.assertNotIn("J2", summary["K4"].value)
        self.assertEqual(cached_summary["K2"].value, 0.5)
        self.assertEqual(cached_summary["K4"].value, 2.0)
        self.assertIn("L2*N2-M2", details["P2"].value)
        self.assertEqual(cached_details["P2"].value, 960)
        self.assertEqual(cached_details["F2"].value, first["images"]["pho"]["sha256"])
        self.assertEqual(cached_details["J2"].value, "8,12,8,6")

    def test_unconfirmed_xlsx_caches_are_blank_and_csv_missing_values_not_zero(self):
        exp = self.synthetic(confirm=False)
        xlsx, _, _ = exports.build_export(self.store, [exp], "xlsx")
        formulas = load_workbook(io.BytesIO(xlsx), data_only=False)["定量结果"]
        cached = load_workbook(io.BytesIO(xlsx), data_only=True)["定量结果"]
        self.assertIsNotNone(formulas["J2"].value)
        self.assertIsNone(cached["J2"].value)
        self.assertIsNone(cached["K2"].value)
        self.assertIsNone(cached["L2"].value)
        self.assertEqual(cached["N2"].value, 0)
        payload, _, _ = exports.build_export(self.store, [exp], "csv")
        row = next(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
        self.assertEqual(row["p/total图像信号比"], "")
        self.assertEqual(row["相对对照图像指标"], "")
        self.assertEqual(float(row["P净信号"]), 960)

    def test_csv_has_independent_provenance_and_treats_sample_names_as_text(self):
        exp = self.synthetic()
        incoming = copy.deepcopy(exp)
        incoming["samples"][0]["name"] = "=2+3"
        self.store.update(exp, incoming)
        payload, mime, extension = exports.build_export(self.store, [exp], "csv")
        self.assertEqual(extension, "csv")
        self.assertTrue(payload.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 2)
        row = rows[0]
        self.assertEqual(row["样本名称"], "'=2+3")
        self.assertEqual(row["实验编号"], exp["id"])
        self.assertEqual(row["样本编号"], exp["samples"][0]["id"])
        self.assertEqual(row["P_SHA256"], exp["images"]["pho"]["sha256"])
        self.assertEqual(row["P_原文件名"], "pho.png")
        self.assertEqual(row["P_条带框xywh"], "8,12,8,6")
        self.assertEqual(row["P_背景框xywh"], "8,2,8,5")
        self.assertEqual(row["P_位深"], "8")
        self.assertIn("未知", row["来源与复核说明"])
        self.assertIn("背景均值", row["计算方法"])

    def test_zip_original_hashes_snapshot_annotations_and_scoped_audit(self):
        exp = self.synthetic("Included")
        other = self.synthetic("Not included")
        original_bytes = {role: self.store.image_path(exp["images"][role]).read_bytes() for role in ("pho", "total")}
        payload, _, _ = exports.build_export(self.store, [exp], "zip")
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = archive.namelist()
            for filename in ("results.xlsx", "results.csv", "provenance.json", "audit.jsonl", "README.txt"):
                self.assertIn(filename, names)
            snapshot = json.loads(archive.read("provenance.json"))
            self.assertEqual([e["id"] for e in snapshot["experiments"]], [exp["id"]])
            self.assertEqual(snapshot["experiments"][0]["rois"], exp["rois"])
            self.assertEqual(snapshot["analyses"][0]["rows"][1]["relative"], 2.0)
            events = [json.loads(line) for line in archive.read("audit.jsonl").decode().splitlines()]
            self.assertTrue(events)
            self.assertTrue(all(event["experimentId"] == exp["id"] for event in events))
            self.assertEqual(events[-1]["revision"], exp["revision"])
            self.assertFalse(any(other["id"] in name for name in names))
            for role, original in original_bytes.items():
                metadata = exp["images"][role]
                path = f"originals/{metadata['id']}/{metadata['filename']}"
                self.assertEqual(archive.read(path), original)
                self.assertEqual(hashlib.sha256(archive.read(path)).hexdigest(), metadata["sha256"])
                image = Image.open(io.BytesIO(archive.read(f"annotated/{exp['id']}-{role}.png")))
                self.assertEqual(image.format, "PNG")
                self.assertGreaterEqual(image.width, metadata["width"])
                self.assertGreater(image.height, metadata["height"])
                self.assertEqual(self.store.image_path(metadata).read_bytes(), original)
            archive_csv = list(csv.DictReader(io.StringIO(archive.read("results.csv").decode("utf-8-sig"))))
            workbook = load_workbook(io.BytesIO(archive.read("results.xlsx")), data_only=True)
            self.assertEqual(float(archive_csv[1]["相对对照图像指标"]), workbook["定量结果"]["L3"].value)


class HTTPBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.store = Store(self.temp.name)
        self.server.verbose = False
        self.thread = threading.Thread(target=self.server.serve_forever, kwargs={"poll_interval": 0.02}, daemon=True)
        self.thread.start()
        self.url = "http://127.0.0.1:" + str(self.server.server_address[1])

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)
        self.temp.cleanup()

    def request(self, route, method="GET", payload=None, headers=None):
        content = json.dumps(payload).encode() if payload is not None else None
        request = Request(self.url + route, data=content, method=method,
                          headers={"Content-Type": "application/json", **(headers or {})})
        try:
            response = urlopen(request, timeout=5)
        except HTTPError as exc:
            response = exc
        with response:
            return response.status, json.loads(response.read())

    def test_health_create_revision_errors_and_cross_origin_rejection(self):
        code, health = self.request("/api/health")
        self.assertEqual(code, 200)
        self.assertTrue(health["ok"])
        code, error = self.request("/api/export?format=xlsx")
        self.assertEqual(code, 400)
        self.assertIn("error", error)
        code, exp = self.request("/api/experiments", "POST", {"name": "HTTP", "sampleCount": 2, "controlCount": 1})
        self.assertEqual(code, 200)
        stale = copy.deepcopy(exp)
        exp["notes"] = "saved"
        code, updated = self.request(f"/api/experiments/{exp['id']}", "PUT", {"experiment": exp})
        self.assertEqual(code, 200)
        self.assertGreater(updated["revision"], stale["revision"])
        code, error = self.request(f"/api/experiments/{exp['id']}", "PUT", {"experiment": stale})
        self.assertEqual(code, 409)
        self.assertIn("error", error)
        bad = copy.deepcopy(updated)
        bad["samples"][1]["id"] = bad["samples"][0]["id"]
        code, error = self.request(f"/api/experiments/{exp['id']}", "PUT", {"experiment": bad})
        self.assertEqual(code, 400)
        code, error = self.request("/api/experiments", "POST", {"name": "Unwanted"}, {"Origin": "https://example.org"})
        self.assertEqual(code, 400)
        code, state = self.request("/api/state")
        self.assertEqual(len(state["experiments"]), 1)
        self.assertEqual(state["experiments"][0]["notes"], "saved")


if __name__ == "__main__":
    unittest.main()
